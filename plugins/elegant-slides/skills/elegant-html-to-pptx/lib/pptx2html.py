#!/usr/bin/env python3
"""PPTX (Régua, W2) -> deck HTML single-file de palco fixo 1920x1080, shape a shape (mesma fonte, tamanho, cor, crop).
Uso: python3 pptx2html.py deck.pptx saida.html"""
import sys, base64, html, pathlib
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

PX = lambda emu: emu / 6350          # 1920 px = 12192000 EMU
BASE = (pathlib.Path(__file__).resolve().parents[2] / 'elegant-html-slides' / 'engine' / 'viewport-base.css').read_text()

def rgb(font, default='141414'):
    try:
        return str(font.color.rgb) if font.color and font.color.type is not None and font.color.rgb is not None else default
    except Exception: return default

def fill_style(sp):
    """cor + alpha de um preenchimento sólido (retângulo/overlay)."""
    sf = sp._element.spPr.find(qn('a:solidFill'))
    if sf is None: return None
    c = sf.find(qn('a:srgbClr'))
    if c is None: return None
    a = c.find(qn('a:alpha')); alpha = int(a.get('val')) / 100000 if a is not None else 1
    r, g, b = int(c.get('val')[0:2], 16), int(c.get('val')[2:4], 16), int(c.get('val')[4:6], 16)
    return f'rgba({r},{g},{b},{alpha:.2f})'

def shape_html(sh):
    x, y, w, h = PX(sh.left), PX(sh.top), PX(sh.width), PX(sh.height)
    pos = f'left:{x:.1f}px;top:{y:.1f}px;width:{w:.1f}px;height:{h:.1f}px;'
    st = sh.shape_type
    if st == 13:  # picture
        blob = sh.image.blob; mime = sh.image.content_type
        cl, cr, ct, cb = sh.crop_left or 0, sh.crop_right or 0, sh.crop_top or 0, sh.crop_bottom or 0
        iw = w / max(1e-6, 1 - cl - cr); ih = h / max(1e-6, 1 - ct - cb)
        src = f'data:{mime};base64,{base64.b64encode(blob).decode()}'
        return (f'<div class="pic" style="{pos}"><img alt="" src="{src}" style="position:absolute;left:{-cl*iw:.1f}px;top:{-ct*ih:.1f}px;width:{iw:.1f}px;height:{ih:.1f}px;max-width:none;max-height:none;"></div>')
    if st == 9:   # line (connector): fio
        col = rgb(sh.line, 'DADAD6') if sh.line.color and sh.line.color.type is not None else 'DADAD6'
        lw = max(1.0, PX(sh.line.width) if sh.line.width else 1.0)
        return f'<div class="rule" style="left:{x:.1f}px;top:{y:.1f}px;width:{max(w, lw):.1f}px;height:{max(h, lw):.1f}px;background:#{col};"></div>'
    if sh.has_text_frame:
        tf = sh.text_frame
        if not tf.text.strip():
            f = fill_style(sh)
            return f'<div style="position:absolute;{pos}background:{f};"></div>' if f else ''
        ml, mt, mr, mb = (PX(v or 0) for v in (tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom))
        anchor = {MSO_ANCHOR.MIDDLE: 'center', MSO_ANCHOR.BOTTOM: 'flex-end'}.get(tf.vertical_anchor, 'flex-start')
        bg = fill_style(sh) if sh.shape_type == 1 else None
        paras = []
        for p in tf.paragraphs:
            al = {PP_ALIGN.CENTER: 'center', PP_ALIGN.RIGHT: 'right', PP_ALIGN.JUSTIFY: 'justify'}.get(p.alignment, 'left')
            ls = p.line_spacing
            lh = 'line-height:1.14;'
            if hasattr(ls, 'pt'): lh = f'line-height:{ls.pt * 2:.1f}px;'
            elif isinstance(ls, float): lh = f'line-height:{ls};'
            sa = f'margin-bottom:{p.space_after.pt * 2:.1f}px;' if p.space_after else ''
            runs = []
            for r in p.runs:
                f = r.font; size = (f.size.pt * 2) if f.size else 36
                name = f.name or 'Inter'
                fam = "'Literata',Georgia,serif" if 'Literata' in name else "'Inter',system-ui,sans-serif"
                wt = 300 if 'Light' in name else (700 if f.bold else 400)
                it = 'italic' if f.italic else 'normal'
                spc = r._r.find(qn('a:rPr')); ls_ = ''
                if spc is not None and spc.get('spc'): ls_ = f'letter-spacing:{int(spc.get("spc")) / 100 * 2:.1f}px;'
                runs.append(f'<span style="font-family:{fam};font-size:{size:.1f}px;font-weight:{wt};font-style:{it};color:#{rgb(f)};{ls_}">{html.escape(r.text)}</span>')
            paras.append(f'<p style="margin:0;text-align:{al};{lh}{sa}">{"".join(runs) or "&nbsp;"}</p>')
        wrap = 'white-space:nowrap;' if tf.word_wrap is False else ''
        bgs = f'background:{bg};' if bg else ''
        return (f'<div class="tb" style="{pos}padding:{mt:.1f}px {mr:.1f}px {mb:.1f}px {ml:.1f}px;justify-content:{anchor};{wrap}{bgs}">{"".join(paras)}</div>')
    f = fill_style(sh)
    return f'<div style="position:absolute;{pos}background:{f};"></div>' if f else ''

def main(src, out):
    prs = Presentation(src)
    title = html.escape(pathlib.Path(src).stem)
    slides = []
    for i, s in enumerate(prs.slides):
        bg = ''
        try:
            if s.background.fill.type == 1: bg = f'background:#{s.background.fill.fore_color.rgb};'
        except Exception: pass
        inner = ''.join(shape_html(sh) for sh in s.shapes)
        slides.append(f'<section class="slide{" active" if i == 0 else ""}" style="{bg}">{inner}</section>')
    doc = f'''<!DOCTYPE html>
<html lang="pt-BR">
<head>
<meta charset="UTF-8"><meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>Antes do PowerPoint, peça para a IA te sabatinar</title>
<link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link href="https://fonts.googleapis.com/css2?family=Literata:ital,opsz,wght@0,7..72,300;0,7..72,400;1,7..72,300;1,7..72,400&family=Inter:wght@400;600;700&display=swap" rel="stylesheet">
<style>
:root{{--stage-bg:#22232A;--slide-bg:#FCFCFC;}}
*{{margin:0;padding:0;box-sizing:border-box;}}
{BASE}
.slide{{background:#FCFCFC;color:#141414;}}
.pic{{position:absolute;overflow:hidden;}}
.rule{{position:absolute;}}
.tb{{position:absolute;display:flex;flex-direction:column;overflow-wrap:break-word;}}
.slide{{transition:opacity .12s;}}
@media (prefers-reduced-motion:reduce){{.slide{{transition:none;}}}}
</style>
</head>
<body>
<div class="deck-viewport"><main class="deck-stage" id="deckStage">
{"".join(slides)}
</main></div>
<script>
(function(){{
  const S=[...document.querySelectorAll('.slide')], stage=document.getElementById('deckStage'); let i=0;
  function show(n){{ i=Math.max(0,Math.min(S.length-1,n)); S.forEach((s,k)=>{{s.classList.toggle('active',k===i);s.classList.toggle('visible',k===i);}}); history.replaceState(null,'','#'+(i+1)); }}
  function fit(){{ const k=Math.min(innerWidth/1920,innerHeight/1080); stage.style.transform=`translate(${{(innerWidth-1920*k)/2}}px,${{(innerHeight-1080*k)/2}}px) scale(${{k}})`; }}
  addEventListener('resize',fit); fit(); show((parseInt(location.hash.slice(1))||1)-1);
  addEventListener('keydown',e=>{{ if(['ArrowRight','ArrowDown',' ','PageDown'].includes(e.key)){{e.preventDefault();show(i+1);}} if(['ArrowLeft','ArrowUp','PageUp'].includes(e.key)){{e.preventDefault();show(i-1);}} if(e.key==='Home')show(0); if(e.key==='End')show(S.length-1); }});
  let x0=null; addEventListener('touchstart',e=>x0=e.touches[0].clientX,{{passive:true}});
  addEventListener('touchend',e=>{{ if(x0===null)return; const d=e.changedTouches[0].clientX-x0; if(Math.abs(d)>40)show(d<0?i+1:i-1); x0=null; }});
  addEventListener('click',e=>{{ show(e.clientX>innerWidth/2?i+1:i-1); }});
  window.deck={{show, next:()=>show(i+1), prev:()=>show(i-1), get index(){{return i;}}, total:S.length}};
}})();
</script>
</body>
</html>'''
    pathlib.Path(out).write_text(doc)
    print(len(prs.slides), 'slides ->', out, f'{len(doc) / 1e6:.1f} MB')

if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
