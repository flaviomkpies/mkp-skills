#!/usr/bin/env python3
"""[genérico] Deck HTML da Régua (/elegant-html-slides) -> PowerPoint NATIVO editável, shape a shape.
Lê a geometria real do DOM (Playwright, palco 1920x1080 sem escala) e reconstrói com pptxlib:
texto vira caixa nativa (Fraunces/Inter, pt-BR), fio vira retângulo de 1px, barra vira shape.
Uso: python3 html2pptx.py deck.html saida.pptx  (gera também saida_TEXTO.md)
"""
import sys, os, re
import pathlib as _pl; sys.path.insert(0, str(_pl.Path(__file__).resolve().parent))
import pptxlib as L
from pptxlib import new_prs, add_slide, add_text, run, E
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from patchright.sync_api import sync_playwright

CHROME = '/opt/google/chrome/chrome'

PROBE = r"""
(idx) => {
  document.querySelectorAll('.slide').forEach((el, k) => { el.classList.toggle('active', k === idx); el.classList.toggle('visible', k === idx); });
  document.querySelector('.deck-stage').style.transform = 'none';
  const s = document.querySelectorAll('.slide')[idx];
  const out = [];
  const px = v => parseFloat(v) || 0;
  const style = el => getComputedStyle(el);
  const isBlock = n => n.nodeType === 1 && !['inline', 'inline-block', 'contents'].includes(style(n).display);
  const ownBox = n => n.nodeType === 1 && (px(style(n).marginLeft) || px(style(n).marginRight) || style(n).float !== 'none');
  const runOf = (node, cs) => ({t: node.textContent, font: cs.fontFamily.split(',')[0].replace(/['"]/g, ''),
      size: px(cs.fontSize), bold: parseInt(cs.fontWeight) >= 600, italic: cs.fontStyle === 'italic', color: cs.color});
  function flush(group, block) {
    const runs = group.map(g => g.run).filter(r => r.t.replace(/ /g, ' ').trim() || / /.test(r.t));
    if (!runs.length) return;
    const rg = document.createRange(); rg.setStartBefore(group[0].node); rg.setEndAfter(group[group.length - 1].node);
    const rects = [...rg.getClientRects()].filter(r => r.width > 0);
    if (!rects.length) return;
    const l = Math.min(...rects.map(r => r.left)), t = Math.min(...rects.map(r => r.top));
    const r_ = Math.max(...rects.map(r => r.right)), b = Math.max(...rects.map(r => r.bottom));
    const bcs = style(block), br = block.getBoundingClientRect();
    const bl = br.left + px(bcs.paddingLeft), bw = br.width - px(bcs.paddingLeft) - px(bcs.paddingRight);
    const first = group[0].node.nodeType === 1 ? group[0].node : block;
    const fl = style(first).float;
    let x, w, align = 'left';
    if (bcs.textAlign === 'center') { x = bl; w = bw; align = 'center'; }
    else if (fl === 'right') { x = r_ - (r_ - l) * 1.06 - 6; w = (r_ - l) * 1.06 + 6; align = 'right'; }
    else { x = l; w = Math.min((r_ - l) * 1.05 + 10, bl + bw - l); }
    out.push({k: 'text', x, y: t, w, h: b - t, runs, align, line: px(bcs.lineHeight) || px(bcs.fontSize) * 1.2,
              li: block.tagName === 'LI'});
  }
  function walk(el) {
    const cs = style(el), r = el.getBoundingClientRect();
    if (cs.visibility === 'hidden' || cs.display === 'none' || r.width === 0) return;
    if (el !== s) {
      if (px(cs.borderTopWidth) && cs.borderTopStyle !== 'none') out.push({k: 'rect', x: r.left, y: r.top, w: r.width, h: 1, fill: cs.borderTopColor});
      if (px(cs.borderLeftWidth) && cs.borderLeftStyle !== 'none') out.push({k: 'rect', x: r.left, y: r.top, w: 1, h: r.height, fill: cs.borderLeftColor});
      const bg = cs.backgroundColor;
      if (bg && !/rgba\(0, 0, 0, 0\)|transparent/.test(bg) && !el.textContent.trim()) out.push({k: 'rect', x: r.left, y: r.top, w: Math.max(1, r.width), h: Math.max(1, r.height), fill: bg});
    }
    let group = [];
    for (const n of el.childNodes) {
      if (n.nodeType === 3) { if (n.textContent.trim() || / /.test(n.textContent)) group.push({node: n, run: runOf(n, cs)}); continue; }
      if (n.nodeType !== 1) continue;
      if (isBlock(n)) { flush(group, el); group = []; walk(n); continue; }
      if (ownBox(n)) { flush(group, el); group = []; flush([{node: n, run: runOf(n, style(n))}], el); continue; }
      group.push({node: n, run: runOf(n, style(n))});
    }
    flush(group, el);
  }
  walk(s);
  return out;
}
"""

def hexc(rgb):
    m = re.findall(r'\d+', rgb)
    return '%02X%02X%02X' % tuple(int(v) for v in m[:3])

def rect(slide, x, y, w, h, fill):
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, E(x), E(y), E(w), E(h))
    sp.fill.solid(); sp.fill.fore_color.rgb = L.C(fill); sp.line.fill.background()
    st = sp._element.find(qn('p:style'))            # sem <p:style>: sem sombra herdada do tema Office
    if st is not None: sp._element.remove(st)
    spPr = sp._element.spPr
    if spPr.find(qn('a:effectLst')) is None: spPr.append(spPr.makeelement(qn('a:effectLst'), {}))
    sp.text_frame.paragraphs[0].text = ''
    return sp

def main(html, out):
    html = os.path.abspath(html)
    prs = new_prs(); texto = []
    with sync_playwright() as pw:
        b = pw.chromium.launch(headless=True, executable_path=CHROME, args=['--no-sandbox'])
        pg = b.new_page(viewport={'width': 1920, 'height': 1080})
        pg.goto('file://' + html); pg.wait_for_timeout(1200)     # fontes web
        n = pg.evaluate("document.querySelectorAll('.slide').length")
        for i in range(n):
            items = pg.evaluate(PROBE, i)
            s = add_slide(prs); texto.append(f'### Slide {i + 1}')
            for it in items:
                if it['k'] == 'rect':
                    rect(s, it['x'], it['y'], it['w'], it['h'], hexc(it['fill']))
                else:
                    runs = [run(r['t'].replace(' ', ' '), font=r['font'], size=r['size'], color=hexc(r['color']),
                                bold=r['bold'], italic=r['italic']) for r in it['runs']]
                    if it['li']: runs.insert(0, run('•  ', font=runs[0]['font'], size=runs[0]['size'], color=runs[0]['color']))
                    align = {'center': PP_ALIGN.CENTER, 'right': PP_ALIGN.RIGHT}.get(it['align'], PP_ALIGN.LEFT)
                    x = it['x'] - (26 if it['li'] else 0); w = it['w'] + (26 if it['li'] else 0)
                    add_text(s, x, it['y'], w, it['h'] + 4, [runs], align=align, line_px=it['line'])
                    texto.append(''.join(r['t'] for r in runs).strip())
        b.close()
    prs.save(out)
    open(os.path.splitext(out)[0] + '_TEXTO.md', 'w').write('\n'.join(texto) + '\n')
    print(f'{n} slides -> {out}')

if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
